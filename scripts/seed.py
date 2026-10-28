"""
Загружает разделы услуг, услуги и FAQ из прайс-листа (.xlsx) в базу.

Ожидаемые листы в файле (как в выгрузке из Creatium):
  - "Разделы услуг": Айди | Раздел | Иконка
  - "Услуги": Айди | Раздел | Услуга | Описание | Стоимость полная |
              Стоимость со скидкой | Измерение | Картинка | Порядок для сайта
  - "FAQ": Вопрос | Ответ

Колонка "Раздел" в листе "Услуги" может содержать несколько id через запятую
(например "2, 0") — услуга показывается сразу в нескольких разделах.

Запуск:
    python -m scripts.seed путь/к/прайсу.xlsx
"""

import asyncio
import sys

import openpyxl
from sqlalchemy import select

from app.database import AsyncSessionLocal, init_models
from app.models import Category, Faq, Service


def _clean_rows(ws):
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):
            yield row


async def seed(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    await init_models()

    async with AsyncSessionLocal() as db:
        # --- Разделы услуг ---
        categories_by_source_id: dict[int, Category] = {}
        for row in _clean_rows(wb["Разделы услуг"]):
            source_id, name, icon = row[0], row[1], row[2]
            if source_id is None or name is None:
                continue
            category = Category(name=name, icon=icon, sort_order=int(source_id))
            db.add(category)
            categories_by_source_id[int(source_id)] = category
        await db.flush()  # чтобы у категорий появились свои id

        # --- Услуги ---
        for row in _clean_rows(wb["Услуги"]):
            padded = list(row) + [None] * max(0, 9 - len(row))
            (
                _source_id,
                category_ids_raw,
                title,
                description,
                price_full,
                price_discount,
                unit,
                image_url,
                sort_order,
            ) = padded[:9]

            if title is None:
                continue

            service = Service(
                title=title,
                description=description or "",
                price_full=int(price_full) if price_full is not None else 0,
                price_discount=int(price_discount) if price_discount is not None else None,
                unit=unit or "",
                image_url=image_url,
                sort_order=int(sort_order) if sort_order is not None else 0,
            )

            for raw_id in str(category_ids_raw or "").split(","):
                raw_id = raw_id.strip()
                if not raw_id:
                    continue
                category = categories_by_source_id.get(int(raw_id))
                if category:
                    service.categories.append(category)

            db.add(service)

        # --- FAQ ---
        if "FAQ" in wb.sheetnames:
            for order, row in enumerate(_clean_rows(wb["FAQ"])):
                question, answer = row[0], row[1]
                if not question:
                    continue
                db.add(Faq(question=question, answer=answer or "", sort_order=order))

        await db.commit()

    print("Готово: разделы, услуги и FAQ загружены.")


async def _has_data() -> bool:
    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Category.id).limit(1))
        return result.first() is not None


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Использование: python -m scripts.seed путь/к/прайсу.xlsx")
        sys.exit(1)
    asyncio.run(seed(sys.argv[1]))
